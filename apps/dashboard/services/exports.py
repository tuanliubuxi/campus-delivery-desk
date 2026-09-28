"""Build read-only Excel workbooks from the unified dashboard filter set."""

from datetime import datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.dashboard.selectors import export_fact_sets

HEADER_FILL = PatternFill("solid", fgColor="2456A6")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_sheet(workbook, title, headers, rows):
    """Write a compact review-friendly worksheet with bounded column widths."""
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        # openpyxl rejects timezone-aware datetimes; export local wall-clock values.
        values = [
            timezone.make_naive(value, timezone.get_current_timezone())
            if isinstance(value, datetime) and timezone.is_aware(value)
            else value
            for value in row
        ]
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 36)


def build_dashboard_workbook(filters):
    """Export all required fact families using exactly the dashboard's order set."""
    facts = export_fact_sets(filters)
    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_sheet(
        workbook,
        "订单明细",
        [
            "固定订单号",
            "动态订单号",
            "日期",
            "业务",
            "来源",
            "Agent",
            "收件归属",
            "配送状态",
            "结算状态",
            "加急",
            "上楼",
            "目的区域",
        ],
        (
            (
                order.fixed_id,
                order.display_id,
                order.sequence_date,
                order.get_business_type_display(),
                order.get_source_type_display(),
                order.proxy_batch.agent.name if order.proxy_batch_id else "",
                order.recipient_name_snapshot,
                order.get_delivery_status_display(),
                order.get_settlement_status_display(),
                "是" if order.is_urgent else "否",
                "是" if order.requires_upstairs else "否",
                order.zone_snapshot or order.off_campus_address,
            )
            for order in facts["orders"]
        ),
    )
    _write_sheet(
        workbook,
        "费用项明细",
        ["ID", "订单号", "结算ID", "费用类型", "名称", "数量", "单价", "金额", "状态", "收益人"],
        (
            (
                item.pk,
                item.order.fixed_id if item.order_id else "",
                item.settlement_id or "",
                item.get_charge_type_display(),
                item.label,
                item.quantity,
                item.unit_price,
                item.amount,
                item.get_status_display(),
                str(item.beneficiary_courier or ""),
            )
            for item in facts["charges"]
        ),
    )
    _write_sheet(
        workbook,
        "SettlementLine",
        ["ID", "结算ID", "订单号", "费用类型", "名称", "数量", "单价", "金额", "收益人"],
        (
            (
                line.pk,
                line.settlement_id,
                line.order.fixed_id if line.order_id else "",
                line.get_charge_type_display(),
                line.label,
                line.quantity,
                line.unit_price,
                line.amount,
                str(line.beneficiary_courier or ""),
            )
            for line in facts["lines"]
        ),
    )
    _write_sheet(
        workbook,
        "结算",
        ["ID", "业务", "结算方", "Agent", "代理批次", "状态", "冻结金额", "创建时间", "结算时间"],
        (
            (
                settlement.pk,
                settlement.get_business_type_display(),
                str(settlement.customer or settlement.agent or ""),
                settlement.agent.name if settlement.agent_id else "",
                settlement.proxy_batch.display_code if settlement.proxy_batch_id else "",
                settlement.get_status_display(),
                settlement.amount_due_snapshot,
                settlement.created_at,
                settlement.settled_at,
            )
            for settlement in facts["settlements"]
        ),
    )
    _write_sheet(
        workbook,
        "退款与调整",
        ["ID", "结算ID", "类型", "金额", "原因", "影响工资", "工资成员", "工资金额", "创建时间"],
        (
            (
                adjustment.pk,
                adjustment.settlement_id,
                adjustment.get_adjustment_type_display(),
                adjustment.amount,
                adjustment.reason,
                "是" if adjustment.impact_wage else "否",
                str(adjustment.wage_courier or ""),
                adjustment.wage_amount,
                adjustment.created_at,
            )
            for adjustment in facts["adjustments"]
        ),
    )
    _write_sheet(
        workbook,
        "配送员收益",
        ["ID", "订单号", "结算ID", "配送员", "来源", "收益基数", "比例", "建议工资", "状态"],
        (
            (
                earning.pk,
                earning.order.fixed_id if earning.order_id else "",
                earning.settlement_id or "",
                str(earning.courier),
                earning.get_source_type_display(),
                earning.amount_base,
                earning.commission_rate_snapshot,
                earning.suggested_wage_amount,
                earning.get_status_display(),
            )
            for earning in facts["earnings"]
        ),
    )
    _write_sheet(
        workbook,
        "代理维度",
        ["订单号", "Agent", "批次", "临时收件人", "批次状态"],
        (
            (
                order.fixed_id,
                order.proxy_batch.agent.name,
                order.proxy_batch.display_code,
                order.proxy_recipient.display_name,
                order.proxy_batch.get_status_display(),
            )
            for order in facts["orders"]
            if order.proxy_batch_id
        ),
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
