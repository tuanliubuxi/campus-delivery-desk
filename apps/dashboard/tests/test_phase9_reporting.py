"""Phase 9 acceptance coverage for unified reporting, search, export, and privacy."""

import uuid
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.core.exceptions import PermissionDenied
from django.test import RequestFactory
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.agents.models import Agent, ProxyBatch, ProxyRecipient
from apps.common.enums import BusinessType, UserRole
from apps.customers.models import Customer
from apps.dashboard.selectors import DashboardFilters, dashboard_cards, filtered_orders
from apps.dashboard.services import build_dashboard_workbook
from apps.dashboard.views import courier_statistics, global_search, reports
from apps.dispatch.models import DeliveryDrop, DeliveryDropItem, LocationType
from apps.orders.models import (
    DeliveryStatus,
    DestinationType,
    DispatchMode,
    ExpressOrderDetail,
    ExpressRound,
    Order,
    OrderSettlementStatus,
    PickupArea,
    PickupIdentifierType,
    RecipientKind,
    SizeClass,
    SourceType,
)
from apps.settlements.models import (
    BeneficiaryType,
    ChargeItem,
    ChargeScope,
    ChargeSource,
    ChargeStatus,
    ChargeType,
    CourierEarning,
    EarningSourceType,
    EarningStatus,
    Settlement,
    SettlementLine,
    SettlementOrder,
    SettlementPartyType,
    SettlementStatus,
)


@pytest.fixture
def reporting_facts():
    admin = User.objects.create_user(
        username="phase9-admin",
        display_name="管理员",
        role=UserRole.ADMIN,
    )
    courier = User.objects.create_user(
        username="phase9-courier",
        display_name="配送甲",
        role=UserRole.COURIER,
    )
    other_courier = User.objects.create_user(
        username="phase9-other",
        display_name="配送乙",
        role=UserRole.COURIER,
    )
    customer = Customer.objects.create(wechat_nickname="筛选客户", floor="6", room="601")
    other_customer = Customer.objects.create(wechat_nickname="其他客户", floor="6", room="602")
    service_date = date(2026, 9, 20)
    express_round = ExpressRound.objects.create(
        recipient_kind=RecipientKind.CUSTOMER,
        customer=customer,
        service_date=service_date,
        round_no=1,
    )
    other_round = ExpressRound.objects.create(
        recipient_kind=RecipientKind.CUSTOMER,
        customer=other_customer,
        service_date=service_date,
        round_no=1,
    )
    order = _create_order(
        customer=customer,
        express_round=express_round,
        sequence=1,
        requires_upstairs=True,
        recipient="筛选客户",
    )
    other_order = _create_order(
        customer=other_customer,
        express_round=other_round,
        sequence=2,
        requires_upstairs=False,
        recipient="其他客户",
    )
    settlement = _attach_financial_facts(order=order, customer=customer, courier=courier)
    _attach_delivery(order=order, customer=customer, courier=courier)
    _attach_delivery(order=other_order, customer=other_customer, courier=other_courier)
    return {
        "admin": admin,
        "courier": courier,
        "other_courier": other_courier,
        "order": order,
        "other_order": other_order,
        "settlement": settlement,
    }


def _create_order(*, customer, express_round, sequence, requires_upstairs, recipient):
    order = Order.objects.create(
        business_type=BusinessType.EXPRESS,
        sequence_date=date(2026, 9, 20),
        daily_sequence=sequence,
        service_date=date(2026, 9, 20),
        source_type=SourceType.DIRECT,
        customer=customer,
        delivery_status=DeliveryStatus.DELIVERED,
        settlement_status=OrderSettlementStatus.SETTLED if sequence == 1 else OrderSettlementStatus.UNSETTLED,
        requires_upstairs=requires_upstairs,
        destination_type=DestinationType.CAMPUS_BUILDING,
        building_snapshot="6号楼",
        zone_snapshot="SOUTH",
        floor_snapshot="6",
        room_snapshot="601",
        recipient_name_snapshot=recipient,
    )
    ExpressOrderDetail.objects.create(
        order=order,
        express_round=express_round,
        pickup_area=PickupArea.SOUTH,
        pickup_identifier_type=PickupIdentifierType.PICKUP_CODE,
        pickup_identifier=f"CODE-{sequence}",
        normalized_pickup_identifier=f"CODE{sequence}",
        size_class=SizeClass.LARGE,
        dispatch_mode=DispatchMode.ROUTE,
        small_price_snapshot=Decimal("2.00"),
        medium_price_snapshot=Decimal("4.00"),
        large_price_snapshot=Decimal("6.00"),
        oversize_price_snapshot=Decimal("8.00"),
    )
    return order


def _attach_delivery(*, order, customer, courier):
    drop = DeliveryDrop.objects.create(
        courier=courier,
        recipient_kind=RecipientKind.CUSTOMER,
        customer=customer,
        business_type=BusinessType.EXPRESS,
        building_snapshot="6号楼",
        location_type=LocationType.ROOM,
        final_location_text="6楼房间",
        operation_id=uuid.uuid4(),
        delivered_at=timezone.now(),
    )
    DeliveryDropItem.objects.create(drop=drop, order=order)


def _attach_financial_facts(*, order, customer, courier):
    charge = ChargeItem.objects.create(
        scope_type=ChargeScope.ORDER,
        order=order,
        customer=customer,
        charge_type=ChargeType.BASE_SERVICE,
        label="快递大件基础费",
        quantity=Decimal("1.00"),
        unit_price=Decimal("6.00"),
        amount=Decimal("6.00"),
        source=ChargeSource.SYSTEM_RULE,
        beneficiary_type=BeneficiaryType.PLATFORM,
        status=ChargeStatus.ACTIVE,
    )
    settlement = Settlement.objects.create(
        business_type=BusinessType.EXPRESS,
        party_type=SettlementPartyType.CUSTOMER,
        customer=customer,
        status=SettlementStatus.SETTLED,
        frozen_at=timezone.now(),
        settled_at=timezone.now(),
        amount_due_snapshot=Decimal("6.00"),
    )
    SettlementOrder.objects.create(settlement=settlement, order=order)
    SettlementLine.objects.create(
        settlement=settlement,
        source_charge_item=charge,
        order=order,
        customer=customer,
        charge_type=charge.charge_type,
        label=charge.label,
        quantity=charge.quantity,
        unit_price=charge.unit_price,
        amount=charge.amount,
        source=charge.source,
        beneficiary_type=BeneficiaryType.PLATFORM,
    )
    CourierEarning.objects.create(
        earning_key=f"phase9:{order.pk}",
        order=order,
        settlement=settlement,
        source_charge_item=charge,
        courier=courier,
        source_type=EarningSourceType.BASE_DELIVERY,
        amount_base=Decimal("6.00"),
        commission_rate_snapshot=Decimal("0.5000"),
        suggested_wage_amount=Decimal("3.00"),
        status=EarningStatus.SETTLED,
    )
    return settlement


@pytest.mark.django_db
def test_unified_filters_use_typed_upstairs_and_courier_truth(reporting_facts):
    facts = reporting_facts
    selected = filtered_orders(
        DashboardFilters(
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 30),
            business_type=BusinessType.EXPRESS,
            courier_id=facts["courier"].pk,
            pickup_area=PickupArea.SOUTH,
            size=SizeClass.LARGE,
            route=DispatchMode.ROUTE,
            upstairs=True,
            charge_type=ChargeType.BASE_SERVICE,
        )
    )
    assert list(selected) == [facts["order"]]
    # Both customers have floor data, but the false business flag remains authoritative.
    assert list(filtered_orders(DashboardFilters(upstairs=False))) == [facts["other_order"]]
    cards = dashboard_cards(DashboardFilters(upstairs=True))
    assert cards["settled_income"] == Decimal("6.00")
    # The dashboard reports owned revenue facts; wage percentage is applied only in the calculator.
    assert cards["courier_earnings"] == Decimal("6.00")


@pytest.mark.django_db
def test_agent_dimension_selects_only_matching_proxy_orders():
    agent = Agent.objects.create(name="校园代理甲")
    other_agent = Agent.objects.create(name="校园代理乙")
    batch = ProxyBatch.objects.create(agent=agent, batch_date=date(2026, 9, 21), sequence=1)
    recipient = ProxyRecipient.objects.create(proxy_batch=batch, display_name="临时收件人")
    express_round = ExpressRound.objects.create(
        recipient_kind=RecipientKind.PROXY_RECIPIENT,
        proxy_recipient=recipient,
        service_date=date(2026, 9, 21),
        round_no=1,
    )
    order = Order.objects.create(
        business_type=BusinessType.EXPRESS,
        sequence_date=date(2026, 9, 21),
        daily_sequence=1,
        service_date=date(2026, 9, 21),
        source_type=SourceType.AGENT,
        proxy_recipient=recipient,
        proxy_batch=batch,
        destination_type=DestinationType.CAMPUS_BUILDING,
        building_snapshot="8号楼",
        zone_snapshot="SOUTH",
        recipient_name_snapshot="临时收件人",
    )
    ExpressOrderDetail.objects.create(
        order=order,
        express_round=express_round,
        pickup_area=PickupArea.NORTH,
        pickup_identifier_type=PickupIdentifierType.WAYBILL,
        pickup_identifier="AGENT-001",
        normalized_pickup_identifier="AGENT001",
        size_class=SizeClass.SMALL,
        dispatch_mode=DispatchMode.DIRECT_CUSTOMER,
        small_price_snapshot=Decimal("2.00"),
        medium_price_snapshot=Decimal("4.00"),
        large_price_snapshot=Decimal("6.00"),
        oversize_price_snapshot=Decimal("8.00"),
    )
    assert list(filtered_orders(DashboardFilters(agent_id=agent.pk))) == [order]
    assert not filtered_orders(DashboardFilters(agent_id=other_agent.pk)).exists()


@pytest.mark.django_db
def test_fixed_and_dynamic_global_search_resolve_same_order(reporting_facts):
    factory = RequestFactory()
    for identifier in (reporting_facts["order"].fixed_id, reporting_facts["order"].display_id):
        request = factory.get("/search/", {"q": identifier})
        request.user = reporting_facts["admin"]
        response = global_search(request)
        assert response.status_code == 200
        assert reporting_facts["order"].fixed_id.encode() in response.content
    token_request = factory.get("/search/", {"q": "不存在/筛选客户"})
    token_request.user = reporting_facts["admin"]
    assert reporting_facts["order"].fixed_id.encode() in global_search(token_request).content


@pytest.mark.django_db
def test_excel_inherits_filters_and_contains_required_fact_sheets(reporting_facts):
    payload = build_dashboard_workbook(DashboardFilters(upstairs=True))
    workbook = load_workbook(BytesIO(payload), read_only=True)
    assert set(workbook.sheetnames) == {
        "订单明细",
        "费用项明细",
        "SettlementLine",
        "结算",
        "退款与调整",
        "配送员收益",
        "代理维度",
    }
    order_values = list(workbook["订单明细"].values)
    assert reporting_facts["order"].fixed_id in order_values[1]
    assert all(reporting_facts["other_order"].fixed_id not in row for row in order_values)


@pytest.mark.django_db
def test_role_boundaries_hide_team_reports_and_prevent_courier_spoofing(reporting_facts):
    factory = RequestFactory()
    denied = factory.get("/admin-console/reports/")
    denied.user = reporting_facts["courier"]
    with pytest.raises(PermissionDenied):
        reports(denied)

    request = factory.get(
        "/courier/statistics/",
        {"courier": reporting_facts["other_courier"].pk},
    )
    request.user = reporting_facts["courier"]
    response = courier_statistics(request)
    assert response.status_code == 200
    assert "筛选客户".encode() in response.content
    assert "其他客户".encode() not in response.content
